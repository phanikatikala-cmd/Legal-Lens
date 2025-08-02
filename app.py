import streamlit as st
import re
from datetime import datetime
import PyPDF2
from docx import Document
import requests
from fpdf import FPDF
from transformers import pipeline

# ---- CONFIG ---- #
st.set_page_config(page_title="ClauseWise - AI Legal Document Analyzer", page_icon="⚖️", layout="wide")

# ---- STYLE OVERRIDES ---- #
st.markdown("""
    <style>
    .main {
        background-color: #1E1E2F;
    }
    section[data-testid="stSidebar"] {
        background-color: #0d6efd;
        color: white;
    }
    section[data-testid="stSidebar"] * {
        color: white !important;
    }
    .stFileUploader {
        background-color: #ffffff22;
        border-radius: 10px;
        padding: 8px;
    }
    </style>
""", unsafe_allow_html=True)

# ---- CONSTANTS ---- #
HF_API_KEY = "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"
MODEL_NAME = "ibm-granite/granite-13b-instruct"

# ---- UTILITY FUNCTIONS ---- #
def query_huggingface(prompt, max_tokens=150):
    headers = {"Authorization": f"Bearer {HF_API_KEY}"}
    payload = {
        "inputs": prompt,
        "parameters": {"max_new_tokens": max_tokens, "temperature": 0.3, "return_full_text": False}
    }
    url = f"https://api-inference.huggingface.co/models/{MODEL_NAME}"
    try:
        response = requests.post(url, headers=headers, json=payload)
        result = response.json()
        if isinstance(result, list):
            return result[0]['generated_text'].strip()
        return result.get('generated_text', '')
    except Exception as e:
        return f"Error: {e}"

def extract_text(uploaded_file):
    file_type = uploaded_file.type
    if file_type == "application/pdf":
        reader = PyPDF2.PdfReader(uploaded_file)
        return "\n".join(page.extract_text() for page in reader.pages)
    elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(uploaded_file)
        return "\n".join(para.text for para in doc.paragraphs)
    elif file_type == "text/plain":
        return str(uploaded_file.read(), "utf-8")
    else:
        return "Unsupported file format."

def classify_document(text):
    prompt = f"""Analyze the type of this legal document:\n{text[:500]}\nChoose from:\n1. NDA\n2. Lease\n3. Employment\n4. Service\n5. Purchase\n6. Loan\n7. Other\nRespond with type only."""
    return query_huggingface(prompt, max_tokens=50)

def extract_clauses(text):
    return [cl.strip() for cl in re.split(r'\n\d+\.\s+', text) if len(cl.strip()) > 50][:5]

def simplify_clause(clause):
    prompt = f"""Simplify this clause into plain English:\n\"{clause}\"\nSimplified:"""
    return query_huggingface(prompt, max_tokens=150)

def assess_clause_risk(clause):
    risk_labels = ["Risky", "Neutral", "Safe"]
    classifier = pipeline("zero-shot-classification", model="facebook/bart-large-mnli")
    result = classifier(clause, candidate_labels=risk_labels)
    return result["labels"][0], round(result["scores"][0], 2)

def extract_entities(text):
    ner_pipeline = pipeline("ner", model="dslim/bert-base-NER", aggregation_strategy="simple")
    entities = ner_pipeline(text[:1000])
    return [{"entity": ent['entity_group'], "word": ent['word'], "score": round(ent['score'], 2)} for ent in entities]

def export_pdf(clauses, simplified, entities):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="ClauseWise Report", ln=True, align='C')
    pdf.cell(200, 10, txt="\nEntities:", ln=True)
    for e in entities:
        pdf.cell(200, 10, txt=f"{e['entity']}: {e['word']} ({e['score']})", ln=True)
    pdf.cell(200, 10, txt="\nClauses:", ln=True)
    for i, (orig, simp) in enumerate(zip(clauses, simplified)):
        pdf.multi_cell(0, 10, txt=f"Clause {i+1}:\nOriginal: {orig}\nSimplified: {simp}\n")
    pdf.output("ClauseWise_Report.pdf")

# ---- INTERFACE ---- #
st.title("⚖️ ClauseWise")
st.caption("AI-powered legal document analyzer using IBM Granite + HuggingFace")

with st.sidebar:
    st.image("./A_screenshot_of_a_web_application_named_ClauseWise.png", width=300)
    st.header("📄 Upload Document")
    uploaded_file = st.file_uploader("Choose PDF, DOCX, or TXT", type=["pdf", "docx", "txt"])
    st.markdown("---")
    st.markdown("**Recent Uploads (Coming Soon)**")

if uploaded_file:
    st.success("✅ File uploaded successfully")
    with st.spinner("Extracting text..."):
        raw_text = extract_text(uploaded_file)
        st.session_state.text = raw_text

if 'text' in st.session_state:
    text = st.session_state.text
    st.subheader("📑 Document Preview")
    st.text_area("Preview (First 2000 chars)", text[:2000], height=400)

    tabs = st.tabs(["📋 Classification", "🧠 Entities", "📝 Simplification & Risk", "✅ Final Safety Check", "💬 Clarify Doubts", "📤 Export"])

    with tabs[0]:
        st.subheader("📋 Document Classification")
        doc_type = classify_document(text)
        st.markdown(f"""
            <div style="padding:15px; border-radius:10px;">
                <strong>📄 Predicted Document Type:</strong> <span style="color:orange; font-size:18px;">{doc_type}</span><br>
                <em>Analyzed using the first 500 characters of the document.</em>
            </div>
        """, unsafe_allow_html=True)

    with tabs[1]:
        st.subheader("🧠 Named Entity Recognition")
        entities = extract_entities(text)
        for e in entities:
            st.markdown(f"- **{e['entity']}**: {e['word']} (_score: {e['score']}_)")
        st.session_state.entities = entities

    with tabs[2]:
        st.subheader("📝 Clause Simplification & Risk Analysis")
        clauses = extract_clauses(text)
        simplified = []
        risks = []
        for i, clause in enumerate(clauses):
            with st.expander(f"Clause {i+1}"):
                st.markdown(f"**Original Clause:**\n{clause}")
                simple = simplify_clause(clause)
                risk, score = assess_clause_risk(clause)
                st.success(f"**Simplified:** {simple}")
                st.warning(f"**Risk Level:** {risk} _(Confidence: {score})_")
                simplified.append(simple)
                risks.append(risk)
        st.session_state.clauses = clauses
        st.session_state.simplified = simplified
        st.session_state.risks = risks

    with tabs[3]:
        st.subheader("✅ Final Safety Check")
        if 'risks' in st.session_state:
            if all(r == "Safe" for r in st.session_state.risks):
                st.success("All clauses assessed as Safe. ✅ Document is considered safe.")
            elif any(r == "Risky" for r in st.session_state.risks):
                st.error("⚠️ Risky clauses found. Please review the document.")
            else:
                st.warning("Mixed risk levels. Proceed with caution.")

    with tabs[4]:
       st.subheader("💬 Chatbot for Doubt Clarification")
       st.markdown("If you have any doubts about the document or its clauses, ask here:")

    # Inline input + button using columns
       col1, col2 = st.columns([4, 1])
       with col1:
         user_question = st.text_input(
            "Your question",
            label_visibility="collapsed",
            placeholder="Type your question here..."
          )
       with col2:
           submit = st.button("Submit")

    # On submit
       if submit:
          if not text:
              st.warning("⚠️ Document text is empty. Please upload or extract a document first.")
          elif not user_question.strip():
              st.warning("⚠️ Please enter a question before submitting.")
          else:
              with st.spinner("🤖 Generating answer..."):
                try:
                    prompt = (
                       f"Based on the following document, answer the question:\n"
                       f"{text[:1500]}\n"
                       f"Question: {user_question}\n"
                       f"Answer:"
                    )
                    answer = query_huggingface(prompt)
                    if answer and answer.strip():
                       st.markdown(f"**🗨️ You asked:** {user_question}")
                       st.success("✅ AI's Response:")
                       st.info(answer.strip())
                    else:
                       st.warning("❗ Sorry, I couldn't generate a reply. Try rephrasing your question.")
                except Exception as e:
                     st.error(f"⚠️ An error occurred while querying the model: {e}")


    with tabs[5]:
        st.subheader("📤 Download Analysis Report")
        if st.button("Generate PDF Report"):
            export_pdf(st.session_state.clauses, st.session_state.simplified, st.session_state.entities)
            with open("ClauseWise_Report.pdf", "rb") as f:
                st.download_button("Download Report", f, file_name="ClauseWise_Report.pdf", mime="application/pdf")
else:
    st.info("⬅️ Upload a legal document to begin analysis.")
